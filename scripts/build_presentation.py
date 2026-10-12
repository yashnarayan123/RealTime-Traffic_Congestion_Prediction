"""Build Project_Presentation_Final.pptx from 'Presentation Template.pptx'.

Edits the teacher template in place (title, agenda kept, literature table
columns kept) and appends new slides in the same visual language
(Times New Roman, institutional red C00000, 13.33x7.5in).
All facts are traceable to the repo; placeholders are explicit.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

RED = RGBColor(0xC0, 0x00, 0x00)
DARK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x1F, 0x38, 0x64)
TEAL = RGBColor(0x2E, 0x74, 0xB5)
GREEN = RGBColor(0x37, 0x64, 0x22)
FONT = "Times New Roman"

prs = Presentation("Presentation Template.pptx")
BLANK = next(l for l in prs.slide_layouts if l.name == "Blank")


def set_text(shape, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, font=FONT):
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    p = shape.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.alignment = align
    return shape


def add_box(slide, left, top, width, height):
    return slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))


def new_slide(title, number):
    s = prs.slides.add_slide(BLANK)
    # accent bar + title
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(0.42), Inches(0.9), Inches(0.07))
    bar.fill.solid(); bar.fill.fore_color.rgb = RED; bar.line.fill.background()
    set_text(add_box(s, 0.9, 0.55, 11.5, 0.85), title, size=30, bold=True, color=RED)
    # footer rule + footer
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.72), Inches(11.53), Inches(0.03))
    rule.fill.solid(); rule.fill.fore_color.rgb = RED; rule.line.fill.background()
    set_text(add_box(s, 0.9, 6.78, 9.5, 0.35),
             "Department of Computer Science and Engineering, School of Computing  |  AMARAVATI CAMPUS",
             size=10, color=GRAY)
    set_text(add_box(s, 11.9, 6.78, 0.6, 0.35), str(number), size=10, color=GRAY, align=PP_ALIGN.RIGHT)
    return s


def bullets(slide, items, left=0.9, top=1.55, width=11.5, size=17, gap=6):
    """items: list of (text, bold_prefix_or_None). Returns bottom used."""
    box = add_box(slide, left, top, width, 4.8)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        text, lead = item if isinstance(item, tuple) else (item, None)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        if lead:
            r = p.add_run(); r.text = "• " + lead + " "
            r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = DARK; r.font.name = FONT
            r2 = p.add_run(); r2.text = text
            r2.font.size = Pt(size); r2.font.color.rgb = DARK; r2.font.name = FONT
        else:
            r = p.add_run(); r.text = "•  " + text
            r.font.size = Pt(size); r.font.color.rgb = DARK; r.font.name = FONT
    return box


# ---------- 1. TITLE SLIDE (edit template objects, keep guide/logo/campus) ----------
s1 = prs.slides[0]
for sh in s1.shapes:
    if not sh.has_text_frame:
        continue
    full = "\n".join(p.text for p in sh.text_frame.paragraphs)
    if full.startswith("Tentative Title"):
        set_text(sh, "Tentative Title: Real-Time Traffic Congestion & Accident Analytics",
                 size=22, bold=True, color=RED)
        sh.height = Inches(1.15)  # two lines fit; next block starts at 5.05in
    elif full.startswith("Batch No"):
        set_text(sh, "Batch No: [To be filled]    |    Names of Batch Members (Reg. No.): [To be filled]",
                 size=16, bold=True, color=DARK)

# ---------- 3. LITERATURE SURVEY (template columns kept; rebuilt for 3 rows) ----------
s3 = prs.slides[2]
for sh in list(s3.shapes):
    if sh.has_table:
        sp = sh._element
        sp.getparent().remove(sp)
s3.shapes.add_table(4, 5, Inches(0.55), Inches(1.95), Inches(12.23), Inches(4.6))
tbl = None
for sh in s3.shapes:
    if sh.has_table:
        tbl = sh.table
headers = ["Reference", "Year of Publication", "Datasets used", "Description",
           "Limitations of existing work / Research Gaps"]
widths = [2.0, 1.3, 2.2, 3.5, 3.23]
for j, h in enumerate(headers):
    tbl.columns[j].width = Inches(widths[j])
rows = [
    ["[1] Y. Li, R. Yu, C. Shahabi, Y. Liu — Diffusion Convolutional Recurrent Neural Network (DCRNN), ICLR.",
     "2018", "METR-LA; PEMS-BAY (325 sensors, Bay Area, Jan–May 2017)",
     "Models traffic as diffusion on a directed graph; GRU encoder–decoder with scheduled sampling; +12–15% over baselines.",
     "Heavy graph-RNN training cost; long-horizon error still grows; needs full sensor graph (not a live-incident view)."],
    ["[2] B. Yu, H. Yin, Z. Zhu — Spatio-Temporal Graph Convolutional Networks (STGCN), IJCAI, pp. 3634–3640.",
     "2018", "BJER4M / PeMS-style loop-detector networks",
     "Fully convolutional ST-Conv blocks (temporal gating + graph conv); faster training, fewer parameters than RNN designs.",
     "Fixed graph structure; limited incident/accident modelling; mid/long-term accuracy degrades."],
    ["[3] J. Hogue — Metro Interstate Traffic Volume, UCI Machine Learning Repository (DOI: 10.24432/C5X60B).",
     "2019", "I-94 Minneapolis–St Paul, hourly, 2012–2018, 48,204 rows; weather + holiday features",
     "Public regression benchmark linking weather, holidays and time to traffic volume; used here for batch analytics.",
     "Single station, hourly granularity — no per-sensor network, no real-time stream, no accident/congestion labels."],
]
for i, row in enumerate([headers] + rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
        cell.text_frame.word_wrap = True
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(12 if j > 0 else 11)
        p.font.bold = (i == 0)
        p.font.color.rgb = WHITE if i == 0 else DARK
        p.font.name = FONT

N = 4  # next slide number (1 title, 2 agenda, 3 lit survey kept)


def agenda_note(s, items):
    bullets(s, items, size=18, gap=8)


# ---------- 4. ABSTRACT ----------
s = new_slide("Abstract", N); N += 1
bullets(s, [
    ("Urban congestion and accidents demand both live situational awareness and short-horizon forecasts.", None),
    ("This project builds a streaming analytics pipeline: a Python producer emits traffic events to Kafka; Spark Structured Streaming aggregates them into Parquet; a Streamlit dashboard visualises congestion, accidents and trends.", None),
    ("Separately, RandomForest regressors trained on sampled PEMS-BAY data forecast traffic at 5-, 15- and 60-minute horizons (recorded test R²: 0.966 / 0.889 / 0.600).", None),
    ("The live pipeline is verified end-to-end (Kafka offsets, Spark micro-batches, Parquet output, dashboard); 16/16 project tests pass.", None),
])

# ---------- 5. INTRODUCTION ----------
s = new_slide("Introduction and Background", N); N += 1
bullets(s, [
    ("Traffic data has volume (thousands of sensor events), velocity (continuous streams) and variety (events, weather, sensor series) — a classic Big Data problem.", None),
    ("Batch processing explains the past; stream processing (Kafka + Spark Structured Streaming) reacts to the present micro-batch by micro-batch.", None),
    ("Stack: Python producer → Kafka topic traffic-events → Spark aggregation (congestion %, levels, accident counts) → Parquet → Streamlit dashboard + map.", None),
    ("Historical branch: PEMS-BAY sensor windows (12 lag features) → RandomForest models for 5/15/60-min forecasts, shown on a dedicated dashboard page.", None),
])

# ---------- 6. MOTIVATION ----------
s = new_slide("Motivation of the Work", N); N += 1
bullets(s, [
    ("Commuters and traffic managers need to see congestion forming now, not in yesterday's report.", None),
    ("Incidents are rare but costly: accident events must be counted and localised per road and per window.", None),
    ("Short-horizon forecasts (5–60 min) let operators anticipate build-up; our recorded 5-min R² of 0.966 shows near-term predictability.", None),
    ("A reproducible student-scale stack (Docker Compose, local .venv) demonstrates DBMS, batch/streaming and MapReduce ideas on real data.", None),
])

# ---------- 7. OBJECTIVES ----------
s = new_slide("Objectives of the Work", N); N += 1
bullets(s, [
    ("Ingest simulated live traffic events through Kafka with a validated, versioned event schema.", None),
    ("Aggregate streams with Spark Structured Streaming into restartable Parquet output (congestion %, four levels, accident counts).", None),
    ("Serve a Streamlit dashboard that degrades gracefully (empty states, filter awareness) and maps per-road status.", None),
    ("Forecast PEMS-BAY traffic at 5/15/60-min horizons with trained RandomForest models and report RMSE/MAE/R² honestly.", None),
    ("Verify everything: 16/16 tests, live offset→batch→Parquet→dashboard evidence, one-command start/stop.", None),
])

# ---------- 8. METHODOLOGY ----------
s = new_slide("Methodology and Overall Approach", N); N += 1
bullets(s, [
    ("Streaming path: producer (1 event/s, 5 roads, incident injection) → Kafka (3 partitions) → Spark foreachBatch road aggregation every 10 s → Parquet append + checkpointing.", None),
    ("Forecasting path: 12-lag PEMS-BAY windows → per-horizon RandomForest (scikit-learn 1.6.1) → sample predictions + recorded test metrics on the dashboard page.", None),
    ("Batch analytics: schema checks, sensor summaries, throughput stats and a labelled local MapReduce simulation (pandas, not Hadoop).", None),
    ("Engineering: Linux paths, project-local .venv, Docker Compose (Kafka 4.1, Spark 4.0.1 + Kafka connector), scripts/start.sh --full-stack.", None),
])

# ---------- 9. ARCHITECTURE (editable shapes) ----------
s = new_slide("Model / System Architecture", N); N += 1
nodes = [
    ("Python\nProducer", "1 event/s\nvalidated JSON"),
    ("Kafka\ntraffic-events", "3 partitions\nhost :9094 / ctr :9092"),
    ("Spark Structured\nStreaming", "foreachBatch\n10 s trigger"),
    ("Parquet\ntraffic_summary", "per-road\naggregates"),
    ("Streamlit\nDashboard", "KPIs, trends\nmap, forecast"),
]
x0, y0, bw, bh, gap = 0.55, 2.0, 2.05, 1.55, 0.62
fills = [TEAL, RED, NAVY, TEAL, RED]
for i, (title, sub) in enumerate(nodes):
    x = x0 + i * (bw + gap)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y0), Inches(bw), Inches(bh))
    box.fill.solid(); box.fill.fore_color.rgb = fills[i]; box.line.fill.background()
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = title; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = sub; r2.font.size = Pt(12); r2.font.color.rgb = WHITE; r2.font.name = FONT
    if i < len(nodes) - 1:
        cx = x + bw
        conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(cx + 0.05), Inches(y0 + bh / 2),
                                      Inches(cx + gap - 0.05), Inches(y0 + bh / 2))
        conn.line.color.rgb = RED; conn.line.width = Pt(3)
# forecasting branch (dashed container)
fbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.6), Inches(4.15), Inches(8.1), Inches(1.15))
fbox.fill.background(); fbox.line.color.rgb = NAVY; fbox.line.width = Pt(2)
set_text(add_box(s, 2.8, 4.25, 7.7, 0.95),
         "Offline branch:  PEMS-BAY 12-lag windows  →  RandomForest (5 / 15 / 60 min)  →  Forecasting page  (sample predictions + recorded test metrics)",
         size=15, color=NAVY, align=PP_ALIGN.CENTER)
set_text(add_box(s, 0.9, 5.55, 11.5, 0.9),
         "Checkpointing makes the stream restartable (at-least-once); dashboard reads the same Parquet the stream writes.",
         size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ---------- 10. DATASETS ----------
s = new_slide("Details of the Dataset Used", N); N += 1
bullets(s, [
    ("Live-simulated stream: 5 roads (Central Avenue, Market Road, Airport Road, River Bridge, Highway Junction) with speed, vehicle counts and accident flags; 60-event bounded run filled Kafka offsets 15/21/24 across 3 partitions.", "Streaming events:"),
    ("Batch analytics data: Metro Interstate Traffic Volume (UCI, Hogue 2019) — 48,204 hourly rows, I-94 Minneapolis–St Paul 2012–2018, weather + holiday features; never substituted for PEMS-BAY.", "Metro Interstate:"),
    ("Forecasting data: PEMS-BAY Bay-Area sample (325 sensors, ≤150k rows/split, seed 42) with 12 lag features x_t-11…x_t and targets y_t+1/+3/+12. The small test CSV is still to be placed in data/raw/ — the page shows guidance until then.", "PEMS-BAY (historical):"),
    ("Streaming output: Parquet summaries (window, road, avg speed, congestion %, level, vehicles, accidents) — 700+ rows and growing; zero nulls in validation.", "Processed output:"),
])

# ---------- 11. METRICS ----------
s = new_slide("Evaluation Metrics Used", N); N += 1
bullets(s, [
    ("RMSE (root mean square error): penalises large forecast misses; reported in dataset speed units.", None),
    ("MAE (mean absolute error): average absolute miss; robust companion to RMSE.", None),
    ("R² (coefficient of determination): fraction of variance explained; 1.0 is perfect, 0 equals the mean baseline.", None),
    ("Streaming health: Kafka offsets vs consumed batches, Parquet row growth, null/schema checks, dashboard HTTP 200 with zero log exceptions.", None),
    ("Regression safety: 16/16 pytest checks (serialization, aggregation, empty/malformed handling, metric definitions).", None),
])

# ---------- 12. COMPARATIVE RESULTS (genuine recorded test metrics + chart) ----------
s = new_slide("Comparative Results (Recorded Test Metrics)", N); N += 1
set_text(add_box(s, 0.9, 1.35, 11.5, 0.4),
         "Recorded test-split metrics from the sampled PEMS-BAY experiment (not live-computed; test CSV pending). Deployed model: RandomForest.",
         size=13, color=GRAY)
s.shapes.add_table(6, 5, Inches(0.9), Inches(1.9), Inches(6.4), Inches(2.6))
rt = None
for sh in s.shapes:
    if sh.has_table and sh.top == Inches(1.9):
        rt = sh.table
hdr = ["Horizon", "Model", "RMSE ↓", "MAE ↓", "R² ↑"]
data = [
    ["5 min", "Persistence", "1.936", "0.982", "0.960"],
    ["5 min", "HistGradientBoosting", "1.818", "0.937", "0.965"],
    ["5 min", "RandomForest ★", "1.795", "0.927", "0.966"],
    ["15 min", "RandomForest ★", "3.228", "1.531", "0.889"],
    ["60 min", "RandomForest ★", "6.126", "2.965", "0.600"],
]
for i, row in enumerate([hdr] + data):
    for j, val in enumerate(row):
        cell = rt.cell(i, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = val; p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(13); p.font.bold = (i == 0 or "★" in val)
        p.font.color.rgb = WHITE if i == 0 else DARK; p.font.name = FONT
chart_data = ChartData()
chart_data.categories = ["5 min", "15 min", "60 min"]
chart_data.add_series("Persistence", (1.936, 3.444, 7.014))
chart_data.add_series("HistGradientBoosting", (1.818, 3.229, 6.132))
chart_data.add_series("RandomForest", (1.795, 3.228, 6.126))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.7), Inches(1.9), Inches(4.7), Inches(3.4), chart_data)
ch = gf.chart
ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.value_axis.has_title = True; ch.value_axis.axis_title.text_frame.text = "Test RMSE ↓"
cat = ch.plots[0]
cat.has_data_labels = False
set_text(add_box(s, 0.9, 5.55, 11.5, 0.9),
         "Reading: RandomForest wins all horizons; error grows with horizon (5-min R² 0.966 → 60-min R² 0.600) — near-term traffic is far more predictable.",
         size=14, color=GRAY, align=PP_ALIGN.CENTER)

# ---------- 13. LIVE VERIFICATION ----------
s = new_slide("Implementation Status: Verified Live", N); N += 1
bullets(s, [
    ("Producer → Kafka verified: 60 bounded events landed at offsets 15/21/24; schemas match Spark's from_json on all 10 fields.", "Streaming:"),
    ("Spark → Parquet verified: micro-batches (e.g. 5 road summaries) appended; output read back with valid coords, 0 nulls; checkpoint resume reprocessed only the failed batch.", "Processing:"),
    ("Dashboard verified: HTTP 200, health ok, zero log exceptions with live data; empty-state, filter-empty and missing-CSV warnings all render.", "Serving:"),
    ("One command runs everything (scripts/start.sh --full-stack, duplicate-safe); scripts/stop.sh stops cleanly without deleting data.", "Operations:"),
    ("Not claimed: browser-pixel rendering; live PEMS-BAY scoring (test CSV absent — sample-level predictions only).", "Limits:"),
])

# ---------- 14. CONCLUSION ----------
s = new_slide("Conclusion", N); N += 1
bullets(s, [
    ("A working Big Data pipeline now runs on Linux: simulated traffic flows from producer to dashboard in ~10-second micro-batches.", None),
    ("Congestion percentage, four congestion levels and accident counts are computed per road and visualised with trends and a live map.", None),
    ("Recorded PEMS-BAY results favour RandomForest at every horizon, with strong short-horizon skill (R² 0.966 at 5 min).", None),
    ("Genuine gaps remain: the PEMS-BAY test CSV must be supplied before live forecast scoring; the map needs a browser check.", None),
])

# ---------- 15. FUTURE SCOPE ----------
s = new_slide("Future Scope", N); N += 1
bullets(s, [
    ("Score and monitor the trained models on the full PEMS-BAY test split once the CSV is available; track drift.", None),
    ("Replace/augment the simulator with real sensor or floating-car feeds; add incident alerts.", None),
    ("Explore graph-based forecasters (cf. DCRNN/STGCN in the survey) against the RandomForest baseline.", None),
    ("Harden operations: alerting on consumer lag, Parquet retention policy, dashboard authentication.", None),
])

# ---------- 16. REFERENCES (verified) ----------
s = new_slide("References", N); N += 1
bullets(s, [
    ("Y. Li, R. Yu, C. Shahabi, Y. Liu, “Diffusion Convolutional Recurrent Neural Network: Data-Driven Traffic Forecasting,” Proc. ICLR, 2018. (PEMS-BAY: 325 sensors, Bay Area, Jan–May 2017.)", "[1]"),
    ("B. Yu, H. Yin, Z. Zhu, “Spatio-Temporal Graph Convolutional Networks: A Deep Learning Framework for Traffic Forecasting,” Proc. IJCAI, pp. 3634–3640, 2018.", "[2]"),
    ("J. Hogue, “Metro Interstate Traffic Volume” [Dataset], UCI Machine Learning Repository, 2019. DOI: 10.24432/C5X60B. (I-94, 48,204 hourly rows, 2012–2018.)", "[3]"),
    ("Project artefacts: models/experiment_metadata.json (sample spec, seed 42); models/model_comparison.csv (recorded metrics); batch_experiment/data/raw/Metro_Interstate_Traffic_Volume.csv.", "[4]"),
], size=15, gap=10)

# ---------- 17. THANK YOU ----------
s = new_slide("Thank You", N); N += 1
set_text(add_box(s, 0.9, 2.4, 11.5, 1.2), "Thank You", size=54, bold=True, color=RED, align=PP_ALIGN.CENTER)
set_text(add_box(s, 0.9, 3.7, 11.5, 0.6), "Questions and discussion are welcome.",
         size=20, color=GRAY, align=PP_ALIGN.CENTER)
set_text(add_box(s, 0.9, 4.6, 11.5, 0.6), "Guide: Dr. V Lakshmi Chetana, Assistant Professor (Sl. Gr.), Dept. of CSE",
         size=16, color=DARK, align=PP_ALIGN.CENTER)

prs.save("Project_Presentation_Final.pptx")
print("saved Project_Presentation_Final.pptx with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
